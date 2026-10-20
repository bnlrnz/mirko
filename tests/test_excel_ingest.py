"""Tests for the pre-meeting allocation workbook parser and ingestion."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook

from mirko.ftp import FakeFtpClient
from mirko.ingest.agenda_ingest import AgendaIngestor
from mirko.ingest.excel_ingest import parse_allocation_excel
from mirko.meeting_folder import MeetingFolder
from mirko.worker import SyncWorker

NOW = datetime(2026, 10, 12, 18, 0, tzinfo=timezone.utc)

HEADER = [
    "AI",
    "Topic",
    "TDoc #",
    "Title",
    "Source",
    "Type",
    "For",
    "Avail",
    "Treated",
    "Decision",
    "Wdrn",
    "Replaced by",
    "Replaces",
]


def minimal_config() -> dict:
    return {
        "schema_version": 1,
        "meeting": {
            "id": "SA3#130",
            "title": "SA3#130 Prague",
            "timezone": "Europe/Prague",
            "days": ["2026-10-12", "2026-10-16"],
            "hours": {"start": "09:00", "end": "18:00"},
            "tdoc_series_prefix": "S3-26",
        },
        "source": {
            "kind": "local",
            "protocol": "ftp",
            "host": "10.10.10.10",
            "port": 21,
            "base_path": "/Meetings_3GPP_SYNC/SA3",
            "username": "anonymous",
        },
    }


def workbook_bytes(rows: list[list[str]], header: list[str] | None = HEADER) -> bytes:
    book = Workbook()
    sheet = book.active
    if header is not None:
        sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def make_row(tdoc: str, **overrides: str) -> list[str]:
    row: dict[str, str] = {
        "agenda": "5.3",
        "topic": "Topic",
        "tdoc": tdoc,
        "title": f"Title for {tdoc}",
        "source": "Source",
        "type": "CR",
        "for_action": "Approval",
        "avail": "Yes",
        "treated": "No",
        "decision": "",
        "wdrn": "No",
        "replaced_by": "",
        "replaces": "",
    }
    row.update(overrides)
    return [row[key] for key in (
        "agenda",
        "topic",
        "tdoc",
        "title",
        "source",
        "type",
        "for_action",
        "avail",
        "treated",
        "decision",
        "wdrn",
        "replaced_by",
        "replaces",
    )]


def test_parse_allocation_excel_maps_aliased_headers() -> None:
    data = workbook_bytes(
        [
            make_row("S3-264202", decision="revised", replaced_by="S3-264203"),
            make_row("S3-264203", replaces="S3-264202"),
        ]
    )

    result = parse_allocation_excel(data)

    assert [row.tdoc for row in result.rows] == ["S3-264202", "S3-264203"]
    first = result.rows[0]
    assert first.agenda_item == "5.3"
    assert first.title == "Title for S3-264202"
    assert first.type == "CR"
    assert first.for_action == "Approval"
    assert first.avail is True
    assert first.treated is False
    assert first.decision == "revised"
    assert len(result.edges) == 1
    assert result.edges[0].source == "S3-264202"
    assert result.edges[0].target == "S3-264203"


def test_parse_allocation_excel_skips_blank_and_heading_rows() -> None:
    data = workbook_bytes(
        [
            ["", "", "", "SA3 #130 allocation", "", "", "", "", "", "", "", "", ""],
            make_row("S3-264202"),
            ["", "", "", "", "", "", "", "", "", "", "", "", ""],
        ]
    )

    result = parse_allocation_excel(data)

    assert [row.tdoc for row in result.rows] == ["S3-264202"]


def test_parse_allocation_excel_raises_without_header() -> None:
    data = workbook_bytes([["just", "some", "data"]], header=None)

    try:
        parse_allocation_excel(data)
    except ValueError as exc:
        assert "header" in str(exc)
    else:  # pragma: no cover - defensive
        raise AssertionError("expected ValueError")


def test_excel_ingest_fills_gaps_and_records_prep_source(tmp_path: Path) -> None:
    folder = MeetingFolder.create(tmp_path / "SA3-130", minimal_config())
    data = workbook_bytes([make_row("S3-264202")])

    result = AgendaIngestor(folder, clock=lambda: NOW).ingest_excel(
        data, path="S3-26_meeting_allocation.xlsx"
    )

    assert result.status == "ingested"
    connection = folder.connect()
    try:
        document = connection.execute(
            "SELECT decision, state_source, is_provisional, is_tracked "
            "FROM tdoc WHERE tdoc_number = ?",
            ("S3-264202",),
        ).fetchone()
        assert document["decision"] == "open"
        assert document["state_source"] == "inferred"
        assert document["is_provisional"] == 0
        assert document["is_tracked"] == 0

        assertion = connection.execute(
            "SELECT source, verified, active FROM assertion "
            "WHERE predicate = 'decision' AND subject_ref = ?",
            ("S3-264202",),
        ).fetchone()
        assert assertion["source"] == "excel"
        assert assertion["verified"] == 0
        assert assertion["active"] == 1
    finally:
        connection.close()


def test_html_supersedes_excel_but_excel_never_overwrites_html(
    tmp_path: Path,
) -> None:
    folder = MeetingFolder.create(tmp_path / "SA3-130", minimal_config())
    ingestor = AgendaIngestor(folder, clock=lambda: NOW)
    excel = workbook_bytes([make_row("S3-264202", decision="not pursued")])
    ingestor.ingest_excel(excel)

    agreed = workbook_bytes([make_row("S3-264202", decision="revised")])
    parsed = parse_allocation_excel(agreed)
    row = parsed.rows[0]
    html = (
        "<table><tr><th>Agenda</th><th>Topic</th><th>TDoc</th><th>Title</th>"
        "<th>Source</th><th>Type</th><th>For</th><th>Avail</th><th>Treated</th>"
        "<th>Decision</th><th>Wdrn</th><th>Replaced-by</th><th>Replaces</th></tr>"
        f"<tr><td>{row.agenda_item}</td><td>{row.topic}</td><td>{row.tdoc}</td>"
        f"<td>{row.title}</td><td>{row.source}</td><td>{row.type}</td>"
        f"<td>{row.for_action}</td><td>Yes</td><td>No</td>"
        f"<td>{row.decision}</td><td>No</td><td></td><td></td></tr></table>"
    )
    ingestor.ingest(html)

    connection = folder.connect()
    try:
        document = connection.execute(
            "SELECT decision, state_source FROM tdoc WHERE tdoc_number = ?",
            ("S3-264202",),
        ).fetchone()
        assert document["decision"] == "revised"
        assert document["state_source"] == "html"

        excel_assertion = connection.execute(
            "SELECT active, muted FROM assertion "
            "WHERE predicate = 'decision' AND subject_ref = ? "
            "AND source = 'excel'",
            ("S3-264202",),
        ).fetchone()
        assert excel_assertion["active"] == 0
        assert excel_assertion["muted"] == 1
    finally:
        connection.close()

    # Re-ingesting the older Excel must not clobber the HTML projection.
    ingestor.ingest_excel(excel)
    connection = folder.connect()
    try:
        document = connection.execute(
            "SELECT decision, state_source FROM tdoc WHERE tdoc_number = ?",
            ("S3-264202",),
        ).fetchone()
        assert document["decision"] == "revised"
        assert document["state_source"] == "html"
    finally:
        connection.close()


def test_sync_worker_ingests_excel_before_agenda(tmp_path: Path) -> None:
    folder = MeetingFolder.create(tmp_path / "SA3-130", minimal_config())
    client = FakeFtpClient(
        {
            "S3-26_meeting_allocation.xlsx": workbook_bytes(
                [make_row("S3-264202")]
            ),
        }
    )

    result = SyncWorker(folder, client=client, clock=lambda: NOW).run()

    assert result.excel == "ingested"
    assert result.agenda is None
    assert result.as_dict()["excel"] == "ingested"

    connection = folder.connect()
    try:
        tdocs = {
            row["tdoc_number"]
            for row in connection.execute("SELECT tdoc_number FROM tdoc")
        }
    finally:
        connection.close()
    assert tdocs == {"S3-264202"}


REAL_HEADER = [
    "TDoc",
    "Title",
    "Source",
    "Type",
    "For",
    "Agenda item",
    "Agenda item description",
    "TDoc Status",
    "Is revision of",
    "Revised to",
]


def test_parse_real_workbook_schema() -> None:
    data = workbook_bytes(
        [
            [
                "S3-264055",
                "A#3 - KI#3.1 update",
                "Example Company",
                "pCR",
                "Approval",
                "5.3.1.3.1",
                "Solutions (new/updates)",
                "withdrawn",
                "S3-263998",
                "S3-264056",
            ]
        ],
        header=REAL_HEADER,
    )

    result = parse_allocation_excel(data)

    assert len(result.rows) == 1
    row = result.rows[0]
    assert row.tdoc == "S3-264055"
    assert row.agenda_item == "5.3.1.3.1"
    assert row.topic == "Solutions (new/updates)"
    assert row.decision == "withdrawn"
    assert row.withdrawn is True
    assert row.replaces == "S3-263998"
    assert row.replaced_by == "S3-264056"
    assert {edge.kind for edge in result.edges} == {"revision_of", "based_on"}


def test_parse_prefers_tdoc_worksheet() -> None:
    book = Workbook()
    parameters = book.active
    parameters.title = "Parameters"
    parameters.append(["Types of Tdocs"])
    sheet = book.create_sheet("TDoc_List")
    sheet.append(REAL_HEADER)
    sheet.append(
        [
            "S3-264055",
            "A#3 - KI#3.1 update",
            "Example Company",
            "pCR",
            "Approval",
            "5.3.1.3.1",
            "Solutions (new/updates)",
            "noted",
            "",
            "",
        ]
    )
    buffer = io.BytesIO()
    book.save(buffer)

    result = parse_allocation_excel(buffer.getvalue())

    assert [row.tdoc for row in result.rows] == ["S3-264055"]
